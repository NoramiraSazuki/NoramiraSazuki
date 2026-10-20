"""Tests for the local pilot, on SYNTHETIC files only (no real documents are used).

    cd Amira && ../.venv/bin/python -m unittest discover -s pilot/tests -t . -v
    RUN_SLOW=1 ...      also runs the 4,579-page long-document test
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import resource
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from pilot import inventory, jobs, routing, store
from pilot.engines import build_engines
from pilot.engines.base import Engine, PageResult
from pilot.exports import export_job
from pilot.pages import PageSpecError, parse_page_spec, to_spec
from pilot.settings import load_settings
from pilot.tests import synthetic as syn
from pilot.worker import Worker

PACKAGE_PARENT = Path(__file__).resolve().parents[2]
PREFIX = "20260925/DOCUMENTS/Seismic Report/"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


SEISMIC_ONLY = {"filter_column": "Sub-folder", "filter_value": "Seismic Report", "path_prefix": PREFIX}


def make_env(base: Path, inventory_cfg: dict | None = SEISMIC_ONLY, source: Path | None = None, **processing) -> Path:
    """Settings + source folder + data root, all inside a temp folder.

    Default: the original single-folder setup (Seismic Report). ``inventory_cfg=None`` = all batches.
    """
    src = source or base / "src"
    src.mkdir(parents=True, exist_ok=True)
    proc = {"batch_pages": 3, "ocr_dpi": 200, "min_free_gb": 0, "ocr_page_timeout": 120, **processing}
    inv = "".join(f"{k} = {json.dumps(v)}\n" for k, v in (inventory_cfg or {}).items())
    toml = (f'[paths]\ndata_root = "{base / "data"}"\ninventory_workbook = "{base / "inv.xlsx"}"\n'
            f'source_root = "{src}"\n[inventory]\n{inv}[titles]\ntitle_workbook = "{base / "dirt.xlsx"}"\n'
            f'[processing]\n' + "".join(f"{k} = {json.dumps(v)}\n" for k, v in proc.items()))
    p = base / "settings.toml"
    p.write_text(toml, encoding="utf-8")
    return p


def inv_row(rel: str, n: int, sub="Seismic Report", **kw) -> dict:
    return {"File Name": rel.rsplit("/", 1)[-1], "#": n, "Sub-folder": sub, "File Path": PREFIX + rel, **kw}


class CountingEngine(Engine):
    """Fake engine: records calls; can fail on chosen pages a number of times."""

    def __init__(self, name="fake", fail_pages=(), fail_times=1, version="1", on_page=None):
        self.name, self.label, self._version = name, name, version
        self.calls: list[int] = []
        self.fail_left = {p: fail_times for p in fail_pages}
        self.on_page = on_page

    @property
    def version(self):
        return self._version

    def extract_page(self, ctx):
        self.calls.append(ctx.page_no)
        if self.on_page:
            self.on_page(ctx.page_no)
        if self.fail_left.get(ctx.page_no, 0) > 0:
            self.fail_left[ctx.page_no] -= 1
            raise RuntimeError("synthetic failure at /some/secret/path.pdf")
        return PageResult(f"FAKE page {ctx.page_no}", [], "none", "n/a", {"chars": 9})


class Base(unittest.TestCase):
    def setUp(self):
        self.base = Path(tempfile.mkdtemp(prefix="pilot_test_"))
        self.settings_file = make_env(self.base)
        self.s = load_settings(self.settings_file)
        self.src = self.s.source_root

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def catalogue(self, files: dict[str, Path], extra_rows=()):
        rows = [inv_row(rel, i) for i, rel in enumerate(files, start=1)] + list(extra_rows)
        syn.make_inventory(self.s.inventory_workbook, rows)
        con = store.init_db(self.s.db_path)
        store.upsert_documents(con, inventory.load(self.s))
        return con

    def doc_id(self, con, rel):
        return con.execute("SELECT id FROM documents WHERE rel_path=?", (rel,)).fetchone()[0]

    def worker(self, extra=None):
        engines = build_engines(self.s)
        engines.update(extra or {})
        return Worker(self.s, engines=engines)


# ---------------------------------------------------------------- inventory
class InventoryTest(Base):
    def test_resolution_cases(self):
        syn.make_text(self.src / "ok.pdf")
        (self.src / "SUB").mkdir()
        syn.make_text(self.src / "SUB" / "nested.pdf")
        syn.make_text(self.src / syn.nfd("café.pdf"))           # NFD on disk, NFC in the inventory
        syn.make_text(self.src / "Case.pdf")                     # inventory says case.pdf
        syn.make_text(self.src / "unlisted.pdf")
        syn.make_text(self.src / "dup.pdf")
        (self.src / "._ok.pdf").write_bytes(b"\0")              # macOS companion: ignored
        outside = self.base / "outside.pdf"
        syn.make_text(outside)
        (self.src / "link.pdf").symlink_to(outside)
        rows = [inv_row("ok.pdf", 1, **{"Size (bytes)": (self.src / "ok.pdf").stat().st_size}),
                inv_row("SUB/nested.pdf", 2), inv_row("café.pdf", 3), inv_row("case.pdf", 4),
                inv_row("missing.pdf", 5), inv_row("../escape.pdf", 6), inv_row("link.pdf", 7),
                inv_row("dup.pdf", 8), inv_row("dup.pdf", 9), inv_row("other.pdf", 10, sub="Well Report"),
                {"File Name": "x.pdf", "#": 11, "Sub-folder": "Seismic Report", "File Path": "OTHER/x.pdf"}]
        syn.make_inventory(self.s.inventory_workbook, rows)
        before = sha(self.s.inventory_workbook)
        res = inventory.load(self.s)
        self.assertEqual(sha(self.s.inventory_workbook), before, "inventory must not be modified")
        self.assertEqual(res.header_row, 3)
        st = {r.rel_path: (r.resolve_status, r.note) for r in res.rows}
        self.assertEqual(len(res.rows), 10, "rows filtered by Sub-folder, not by position")
        self.assertEqual(st["ok.pdf"][0], "OK")
        self.assertEqual(st["SUB/nested.pdf"][0], "OK")
        self.assertEqual(st["café.pdf"][0], "OK")
        self.assertEqual(st["case.pdf"][0], "AMBIGUOUS")
        self.assertEqual(st["missing.pdf"][0], "MISSING")
        self.assertEqual(st[PREFIX + "../escape.pdf"][0] if PREFIX + "../escape.pdf" in st else
                         next(v for k, v in st.items() if "escape" in k)[0], "BAD_PATH")
        self.assertEqual(st["link.pdf"][0], "OUTSIDE_SOURCE")
        self.assertEqual(st["dup.pdf"][0], "AMBIGUOUS")
        self.assertEqual(next(v for k, v in st.items() if k.endswith("x.pdf"))[0], "PREFIX_MISMATCH")
        self.assertIn("unlisted.pdf", res.unlisted)
        self.assertIn("Case.pdf", res.unlisted)
        self.assertFalse(any(u.startswith("._") for u in res.unlisted))
        nested = next(r for r in res.rows if r.rel_path == "SUB/nested.pdf")
        self.assertEqual(Path(nested.resolved_path), self.src / "SUB" / "nested.pdf")

    def test_source_unavailable(self):
        syn.make_inventory(self.s.inventory_workbook, [inv_row("a.pdf", 1)])
        shutil.rmtree(self.src)
        res = inventory.load(self.s)
        self.assertFalse(res.source_available)
        self.assertEqual(res.rows[0].resolve_status, "MISSING")

    def test_missing_columns_reported(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "File Inventory"
        wb.active.append(["Name", "Folder"])
        wb.save(self.s.inventory_workbook)
        with self.assertRaises(inventory.InventoryError):
            inventory.load(self.s)


class TitleTest(Base):
    def make_dirt(self, rows):
        """Synthetic mapping sheet: title block, header on row 5, titles in column H under another header."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "DOC-Seismic Report"
        ws.append(["DIRt - synthetic"])
        ws.append([None]); ws.append(["BASE"]); ws.append([None, "M"])
        header = ["#", "*File Name*", "File Size", "*File Format*", "*Title*", "*Dataset Type*",
                  "*Sub Dataset Type*", "Data Status", "Description", "Created By", "Created For", "Remarks",
                  "Top Depth", "Top Depth UOM", "Base Depth", "Base Depth UOM", "Spud Date", "Completion Date",
                  "File Location"]
        ws.append(header)
        for i, (rel, title) in enumerate(rows, start=1):
            r = [None] * len(header)
            r[0], r[1], r[4], r[7] = i, rel.rsplit("/", 1)[-1], "CODE", title
            r[18] = "oss://synthetic-bucket/" + PREFIX + rel if rel != "BAD" else "somewhere/else.pdf"
            ws.append(r)
        wb.save(self.s.title_workbook)

    def test_titles_joined_by_path_not_name(self):
        for rel in ("a.pdf", "SUB/a.pdf", "b.pdf", "c.pdf", "d.pdf"):
            (self.src / rel).parent.mkdir(exist_ok=True)
            syn.make_text(self.src / rel, pages=1)
        self.make_dirt([("a.pdf", "Synthetic Survey Report  A"), ("SUB/a.pdf", "Pressure Survey Sub A"),
                        ("b.pdf", None), ("c.pdf", "Dup One"), ("c.pdf", "Dup Two"), ("BAD", "Lost")])
        before = sha(Path(self.s.title_workbook))
        syn.make_inventory(self.s.inventory_workbook, [inv_row(r, i) for i, r in
                                                       enumerate(["a.pdf", "SUB/a.pdf", "b.pdf", "c.pdf", "d.pdf"], 1)])
        from pilot import titles

        t = titles.load(self.s)
        self.assertEqual(sha(Path(self.s.title_workbook)), before, "title workbook must not be modified")
        self.assertEqual((t.header_row, t.title_header), (5, "Data Status"))
        self.assertEqual((t.blank, t.bad_location, t.duplicates), (1, 1, ["c.pdf"]))
        con = store.init_db(self.s.db_path)
        store.upsert_documents(con, inventory.load(self.s), t)
        got = {r[0]: r[1] for r in con.execute("SELECT rel_path, title FROM documents")}
        self.assertEqual(got, {"a.pdf": "Synthetic Survey Report A", "SUB/a.pdf": "Pressure Survey Sub A",
                               "b.pdf": None, "c.pdf": None, "d.pdf": None})
        # a later reload without a title source keeps titles; a missing workbook is reported, not fatal
        store.upsert_documents(con, inventory.load(self.s), None)
        self.assertEqual(con.execute("SELECT title FROM documents WHERE rel_path='a.pdf'").fetchone()[0],
                         "Synthetic Survey Report A")
        Path(self.s.title_workbook).unlink()
        self.assertIn("not found", titles.load(self.s).error)


class AllBatchesTest(unittest.TestCase):
    """Whole share: every PDF in every batch folder, folder tree incl. empty folders."""

    def setUp(self):
        self.base = Path(tempfile.mkdtemp(prefix="pilot_test_"))
        self.share = self.base / "DulangOSS"
        self.seis = self.share / "20260925/DOCUMENTS/Seismic Report"
        (self.seis / "BOTTOM_HOLE_PRESSURE_SURVEY").mkdir(parents=True)
        (self.share / "20260925/CHECKSHOT").mkdir(parents=True)               # no PDFs: still shown
        (self.share / "20261016/DULANG_WELLS/Well_History").mkdir(parents=True)
        syn.make_text(self.seis / "s1.pdf", pages=2)
        syn.make_text(self.seis / "BOTTOM_HOLE_PRESSURE_SURVEY" / "b1.pdf", pages=1)
        syn.make_text(self.share / "20261016/DULANG_WELLS/Well_History/w1.pdf", pages=3)
        (self.share / "20260925/CHECKSHOT/data.las").write_text("~synthetic")
        self.rows = [
            {"File Name": "s1.pdf", "#": 1, "Sub-folder": "Seismic Report", "File Path": PREFIX + "s1.pdf", "Total Pages": 2},
            {"File Name": "b1.pdf", "#": 2, "Sub-folder": "Seismic Report",
             "File Path": PREFIX + "BOTTOM_HOLE_PRESSURE_SURVEY/b1.pdf", "Total Pages": 1},
            {"File Name": "w1.pdf", "#": 3, "Sub-folder": "Well_History",
             "File Path": "20261016/DULANG_WELLS/Well_History/w1.pdf", "Total Pages": 3},
            {"File Name": "m.pdf", "#": 4, "Sub-folder": "01 WIR", "File Path": "20261015/MBAL/01 WIR/m.pdf",
             "Total Pages": 5},
            {"File Name": "data.las", "#": 5, "Sub-folder": "CHECKSHOT", "File Path": "20260925/CHECKSHOT/data.las"},
        ]

    def tearDown(self):
        shutil.rmtree(self.base, ignore_errors=True)

    def settings(self, seismic_only: bool):
        if seismic_only:
            return load_settings(make_env(self.base, source=self.seis))
        return load_settings(make_env(self.base, inventory_cfg=None, source=self.share))

    def test_all_batches_and_folders(self):
        s = self.settings(False)
        syn.make_inventory(s.inventory_workbook, self.rows)
        from pilot import catalogue

        state = catalogue.refresh(s)
        self.assertEqual(state["documents"], 4, "PDF rows only, from every batch")
        con = store.connect(s.db_path)
        got = {r["rel_path"]: r["resolve_status"] for r in con.execute("SELECT rel_path, resolve_status FROM documents")}
        self.assertEqual(got, {"20260925/DOCUMENTS/Seismic Report/s1.pdf": "OK",
                               "20260925/DOCUMENTS/Seismic Report/BOTTOM_HOLE_PRESSURE_SURVEY/b1.pdf": "OK",
                               "20261016/DULANG_WELLS/Well_History/w1.pdf": "OK",
                               "20261015/MBAL/01 WIR/m.pdf": "MISSING"})
        folders = json.loads(store.meta_get(con, "folders"))
        for f in ("20260925", "20260925/CHECKSHOT", "20260925/DOCUMENTS/Seismic Report/BOTTOM_HOLE_PRESSURE_SURVEY",
                  "20261015/MBAL/01 WIR"):  # empty folder, nested folder, inventory-only folder
            self.assertIn(f, folders)

    def test_widening_scope_keeps_history(self):
        s1 = self.settings(True)
        syn.make_inventory(s1.inventory_workbook, self.rows)
        from pilot import catalogue

        catalogue.refresh(s1)
        con = store.connect(s1.db_path)
        did = con.execute("SELECT id FROM documents WHERE rel_path='s1.pdf'").fetchone()[0]
        job = jobs.create_job(con, [(did, "all")])
        w = Worker(s1)
        w.run_once()
        v = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        jobs.add_review(con, v["sha256"], 1, None, "approved")
        s2 = self.settings(False)  # same data root, whole share
        state = catalogue.refresh(s2)
        self.assertEqual(state["rekeyed"], 2)
        row = con.execute("SELECT id, rel_path FROM documents WHERE id=?", (did,)).fetchone()
        self.assertEqual(row["rel_path"], "20260925/DOCUMENTS/Seismic Report/s1.pdf")
        self.assertEqual(con.execute("SELECT COUNT(*) FROM documents WHERE resolved_path=?",
                                     (str(self.seis / "s1.pdf"),)).fetchone()[0], 1, "no duplicate document")
        self.assertEqual(jobs.job_summary(con, job)["items"][0]["doc_id"], did)
        w2 = Worker(s2)
        v2 = jobs.page_view(con, did, 1, w2.engines, w2.rules.version)
        self.assertEqual(v2["review"]["status"], "approved", "review history still attached")
        self.assertEqual(len(list((s2.data_root / "backups").glob("*.sqlite3"))), 1, "database backed up first")

    def test_nas_offline_keeps_catalogue(self):
        s = self.settings(False)
        syn.make_inventory(s.inventory_workbook, self.rows)
        from pilot import catalogue

        catalogue.refresh(s)
        self.share.rename(self.base / "offline")
        state = catalogue.refresh(s)
        self.assertIn("NOT CHECKED", state["error"])
        con = store.connect(s.db_path)
        self.assertEqual(con.execute("SELECT COUNT(*) FROM documents WHERE resolve_status='OK'").fetchone()[0], 3)

    def test_inventory_first_then_progressive_check(self):
        from pilot import catalogue
        from pilot import inventory as inv_mod

        s = self.settings(False)
        syn.make_inventory(s.inventory_workbook, self.rows)
        con = store.init_db(s.db_path)
        store.meta_set(con, "scan_state", json.dumps({"running": True}))  # app closed mid-scan last time
        seen = {}
        real_check = inv_mod.check_on_disk

        def spy(rows, source_root, extensions, progress, folder_timeout, on_checked):
            # before any NAS listing, every inventory PDF is already in the catalogue as NOT_CHECKED
            seen["before"] = {r[0]: r[1] for r in con.execute("SELECT rel_path, resolve_status FROM documents")}
            seen["groups"] = []

            def counting(batch):
                seen["groups"].append(len(batch))
                on_checked(batch)
            return real_check(rows, source_root, extensions, progress, folder_timeout, counting, flush_seconds=0)

        with mock.patch.object(catalogue.inventory, "check_on_disk", spy):
            state = catalogue.refresh(s)
        self.assertEqual(set(seen["before"].values()), {"NOT_CHECKED"})
        self.assertEqual(len(seen["before"]), 4)
        self.assertGreater(len(seen["groups"]), 1, "statuses are written folder by folder, not only at the end")
        self.assertTrue(state["previous_interrupted"])
        self.assertEqual(state["checked"], 4)
        got = {r[0]: r[1] for r in con.execute("SELECT rel_path, resolve_status FROM documents")}
        self.assertEqual(sorted(got.values()), ["MISSING", "OK", "OK", "OK"])
        self.assertIn("20260925/CHECKSHOT", json.loads(store.meta_get(con, "folders")))

    def test_incomplete_listing_is_not_reported_missing(self):
        """The NAS sometimes omits entries from a listing; a second direct look-up must catch that."""
        from pilot import inventory as inv_mod

        s = self.settings(False)
        syn.make_inventory(s.inventory_workbook, self.rows)
        real = inv_mod._scandir_retry
        calls = {"n": 0}

        def partial(d, attempts=3, pause=1.0, timeout=60.0):
            entries = real(d, attempts, 0, timeout)
            if str(d) == str(self.share) and calls["n"] == 0:  # first root listing hides two batch folders
                calls["n"] += 1
                return [e for e in entries if e[0] not in ("20260925", "20261016")]
            return entries

        with mock.patch.object(inv_mod, "_scandir_retry", partial):
            res = inv_mod.resolve(self.rows, self.share, "")
        st = {r.rel_path: r.resolve_status for r in res.rows if r.rel_path.endswith(".pdf")}
        self.assertEqual(st["20260925/DOCUMENTS/Seismic Report/s1.pdf"], "OK")
        self.assertEqual(st["20261016/DULANG_WELLS/Well_History/w1.pdf"], "OK")
        self.assertEqual(st["20261015/MBAL/01 WIR/m.pdf"], "MISSING")
        m = next(r for r in res.rows if r.rel_path.endswith("m.pdf"))
        self.assertIn("confirmed by a second look-up", m.note)

    def test_unreadable_folder_is_reported_not_missing(self):
        from pilot import catalogue
        from pilot import inventory as inv_mod

        s = self.settings(False)
        syn.make_inventory(s.inventory_workbook, self.rows)
        catalogue.refresh(s)
        real = inv_mod._scandir_retry

        def flaky(d, attempts=3, pause=1.0, timeout=60.0):  # the Seismic Report folder keeps failing
            if str(d).endswith("Seismic Report"):
                raise OSError(60, "Operation timed out")
            return real(d, attempts, 0, timeout)

        with mock.patch.object(inv_mod, "_scandir_retry", flaky):
            state = catalogue.refresh(s)
        self.assertEqual(state["failed_folders"], 1)
        self.assertIn("could not be listed", state["error"])
        con = store.connect(s.db_path)
        st = {r[0]: (r[1], r[2]) for r in con.execute("SELECT rel_path, resolve_status, note FROM documents")}
        self.assertEqual(st["20260925/DOCUMENTS/Seismic Report/s1.pdf"][0], "OK", "earlier match kept")
        self.assertIn("previous match kept", st["20260925/DOCUMENTS/Seismic Report/s1.pdf"][1])
        # first-ever scan with a failing folder: FOLDER_UNREADABLE, never MISSING
        shutil.rmtree(s.data_root)
        s.ensure_dirs()
        with mock.patch.object(inv_mod, "_scandir_retry", flaky):
            catalogue.refresh(s)
        con = store.connect(s.db_path)
        self.assertEqual(con.execute("SELECT resolve_status FROM documents WHERE rel_path=?",
                                     ("20260925/DOCUMENTS/Seismic Report/s1.pdf",)).fetchone()[0], "FOLDER_UNREADABLE")

    def test_stalled_folder_listing_is_skipped(self):
        from pilot import inventory as inv_mod

        real_scandir = os.scandir
        release = __import__("threading").Event()

        def stalling(d):  # this folder never answers (like a stalled SMB request)
            if str(d).endswith("Well_History"):
                release.wait(30)
            return real_scandir(d)

        try:
            with mock.patch.object(inv_mod.os, "scandir", stalling):
                t0 = time.time()
                idx = inv_mod._disk_index(self.share, folder_timeout=0.3)
            self.assertLess(time.time() - t0, 10, "scan does not wait for the stalled folder")
            self.assertEqual(idx.failed, ["20261016/DULANG_WELLS/Well_History"])
            self.assertIn("20260925/DOCUMENTS/Seismic Report/s1.pdf", idx.files, "other folders still scanned")
        finally:
            release.set()

    def test_titles_from_all_sheets(self):
        from openpyxl import Workbook

        from pilot import titles

        s = self.settings(False)
        wb = Workbook()
        wb.active.title = "Read Me"
        for name, rows in (("DOC-Seismic Report", [(PREFIX + "s1.pdf", "Seismic Title"),
                                                   (PREFIX + "BOTTOM_HOLE_PRESSURE_SURVEY/b1.pdf", None)]),
                           ("DOC-Well Report", [("20261016/DULANG_WELLS/Well_History/w1.pdf", "Well Title"),
                                                (PREFIX + "s1.pdf", "Listed twice")])):
            ws = wb.create_sheet(name)
            ws.append(["header block"]); ws.append([]); ws.append([]); ws.append([])
            ws.append(["#", "*File Name*", "", "", "*Title*", "", "", "Data Status", "File Location"])
            for loc, title in rows:
                ws.append([1, "x", None, None, "CODE", None, None, title, "oss://synthetic-bucket/" + loc])
        wb.save(s.title_workbook)
        t = titles.load(s)
        self.assertEqual(t.sheets, ["DOC-Seismic Report", "DOC-Well Report"])
        self.assertEqual(t.titles, {"20261016/DULANG_WELLS/Well_History/w1.pdf": "Well Title"},
                         "path listed in two sheets gets no title; blank stays blank")
        self.assertEqual(t.sources["20261016/DULANG_WELLS/Well_History/w1.pdf"], "DOC-Well Report")
        self.assertEqual(t.duplicates, [PREFIX + "s1.pdf"])

    def test_large_job_needs_confirmation(self):
        from fastapi.testclient import TestClient

        from pilot import catalogue
        from pilot.web.app import create_app

        s = self.settings(False)
        s.large_job_pages = 4
        syn.make_inventory(s.inventory_workbook, self.rows)
        catalogue.refresh(s)
        app = create_app(s, allowed_hosts={"testserver"})
        c = TestClient(app)
        h = {"X-Pilot-Token": app.state.token}
        con = store.connect(s.db_path)
        ids = [r[0] for r in con.execute("SELECT id FROM documents WHERE resolve_status='OK' ORDER BY id")]
        items = [{"doc_id": i} for i in ids]  # 2 + 1 + 3 = 6 pages > 4
        r = c.post("/api/jobs", json={"items": items, "mode": "text_metadata"}, headers=h)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.json()["pages"], 6)
        self.assertEqual(c.post("/api/jobs", json={"items": items[:1], "mode": "text_metadata"}, headers=h)
                         .status_code, 200, "small job needs no confirmation")
        r = c.post("/api/jobs", json={"items": items, "mode": "text_metadata", "confirm_large": True}, headers=h)
        self.assertEqual(r.status_code, 200)
        rr = c.post("/api/inventory/reload", headers=h)
        self.assertEqual(rr.status_code, 202)
        for _ in range(300):  # let the background refresh finish before the next test starts one
            if not catalogue._lock.locked():
                break
            time.sleep(0.1)
        app.state.source_monitor.stop()


class PageSpecTest(unittest.TestCase):
    def test_parse(self):
        self.assertEqual(parse_page_spec("1-3, 7,9-10"), [1, 2, 3, 7, 9, 10])
        self.assertEqual(parse_page_spec("all", 3), [1, 2, 3])
        self.assertEqual(to_spec([10, 9, 7, 3, 2, 1]), "1-3,7,9-10")
        for bad in ("0", "5-2", "a", "1-", "3-4x"):
            with self.assertRaises(PageSpecError):
                parse_page_spec(bad)
        with self.assertRaises(PageSpecError):
            parse_page_spec("4-6", 5)


# ---------------------------------------------------------------- routing
class RoutingTest(Base):
    def test_mixed_document_routing(self):
        pdf_path = syn.make_mixed(self.base / "mixed.pdf")
        from pilot.pdfio import open_pdf

        rules = routing.load_rules()
        pdf = open_pdf(pdf_path)
        try:
            got = {n: routing.decide(routing.page_evidence(pdf[n - 1]), rules)[0] for n in range(1, len(pdf) + 1)}
        finally:
            pdf.close()
        self.assertEqual(got, syn.MIXED_EXPECTED)

    def test_bad_rule_rejected(self):
        bad = self.base / "r.json"
        bad.write_text(json.dumps({"rules": [{"when": {"colour_gt": 1}, "method": "tesseract"}]}))
        with self.assertRaises(ValueError):
            routing.load_rules(bad)


# ---------------------------------------------------------------- worker end to end
class WorkerTest(Base):
    def test_mixed_range_page_numbers_and_export(self):
        syn.make_mixed(self.src / "mixed.pdf")
        before = sha(self.src / "mixed.pdf")
        con = self.catalogue({"mixed.pdf": None})
        did = self.doc_id(con, "mixed.pdf")
        job = jobs.create_job(con, [(did, "1-3,5,7")])
        w = self.worker()
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done")
        self.assertEqual((j["pages_total"], j["pages_done"]), (5, 5))
        self.assertEqual(sha(self.src / "mixed.pdf"), before, "source must be unchanged")
        self.assertEqual(sorted(p.name for p in self.src.iterdir()), ["mixed.pdf"], "nothing written beside source")
        # original page numbers survive range selection + batching (batch_pages=3)
        v = jobs.page_view(con, did, 5, w.engines, w.rules.version)
        self.assertEqual(v["selected_method"], "tesseract")
        ocr = next(r for r in v["results"] if r["engine"] == "tesseract")
        self.assertIn("ROTATED SCAN PAGE FIVE", ocr["text"])
        v1 = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        self.assertIn("PAGE-MARKER 1", next(r for r in v1["results"] if r["engine"] == "textlayer")["text"])
        self.assertEqual({r["page_no"] for r in con.execute("SELECT DISTINCT page_no FROM page_results")},
                         {1, 2, 3, 5, 7})
        # figure page: OCR finds less than the digital caption -> flagged, nothing dropped silently
        out = export_job(con, self.s, job, w.engines, w.rules.version)
        self.assertFalse(str(out).startswith(str(PACKAGE_PARENT)), "exports stay outside the repository")
        pages = list(csv.DictReader(open(out / "pages.csv", encoding="utf-8-sig")))
        self.assertEqual([int(p["Page"]) for p in pages], [1, 2, 3, 5, 7])
        p7 = next(p for p in pages if p["Page"] == "7")
        self.assertIn("text_layer_has_more_text_than_selected_result", p7["Validation Flags"])
        txt = (out / f"doc{did}" / "text.txt").read_text(encoding="utf-8")
        self.assertIn("===== Page 5 (method=tesseract", txt)
        meta = list(csv.DictReader(open(out / "metadata.csv", encoding="utf-8-sig")))[0]
        self.assertEqual(meta["Created For"], "Acme Exploration Sdn Bhd")
        self.assertIn("not verified by a reviewer", meta["EXTRACTION STATUS"])
        chunks = [json.loads(l) for l in open(out / f"doc{did}" / "chunks.jsonl", encoding="utf-8")]
        self.assertTrue(all({"page_no", "sha256", "rel_path", "method", "text_origin"} <= set(c) for c in chunks))
        manifest = json.loads((out / "manifest.json").read_text())
        self.assertIn("does not mean the extraction is accurate", manifest["caveat"])
        time.sleep(1.1)  # export folders are timestamped to the second
        out2 = export_job(con, self.s, job, w.engines, w.rules.version)
        self.assertNotEqual(out, out2)
        self.assertTrue((out / "manifest.json").is_file(), "earlier export untouched")

    def test_failed_page_kept_then_retry_reuses_completed(self):
        syn.make_text(self.src / "t.pdf", pages=7)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        fake = CountingEngine("fake", fail_pages=[4])
        w = self.worker({"fake": fake})
        job = jobs.create_job(con, [(did, "all")], method_override="fake")
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done_with_errors")
        self.assertEqual(j["pages_failed"], 1)
        self.assertEqual(sorted(fake.calls), [1, 2, 3, 4, 5, 6, 7])
        err = con.execute("SELECT error FROM page_results WHERE engine='fake' AND page_no=4").fetchone()[0]
        self.assertNotIn("/some/secret", err, "errors are sanitised")
        ok_before = {r["page_no"]: r["id"] for r in con.execute(
            "SELECT page_no, id FROM page_results WHERE engine='fake' AND status='ok'")}
        self.assertEqual(sorted(ok_before), [1, 2, 3, 5, 6, 7])
        fake.calls.clear()
        jobs.retry(con, job)
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done")
        self.assertEqual(fake.calls, [4], "only the failed page is redone")
        ok_after = {r["page_no"]: r["id"] for r in con.execute(
            "SELECT page_no, id FROM page_results WHERE engine='fake' AND status='ok'")}
        self.assertEqual({k: ok_after[k] for k in ok_before}, ok_before)

    def test_pause_resume_cancel(self):
        syn.make_text(self.src / "t.pdf", pages=9)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        job = jobs.create_job(con, [(did, "all")], method_override="fake")
        ui = store.connect(self.s.db_path)
        fake = CountingEngine("fake", on_page=lambda n: n == 4 and jobs.pause(ui, job))
        w = self.worker({"fake": fake})
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "paused")
        self.assertEqual(fake.calls, [1, 2, 3, 4])
        jobs.resume(con, job)
        fake.on_page = lambda n: n == 6 and jobs.cancel(ui, job)
        w.run_once()
        self.assertEqual(jobs.job_summary(con, job)["status"], "cancelled")
        self.assertEqual(fake.calls, [1, 2, 3, 4, 5, 6])
        fake.on_page = None
        jobs.resume(con, job)
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done")
        self.assertEqual(fake.calls, list(range(1, 10)), "each page processed exactly once")
        n = con.execute("SELECT COUNT(*) FROM page_results WHERE engine='fake'").fetchone()[0]
        self.assertEqual(n, 9)

    def test_changed_source_not_reused_and_settings_hash(self):
        syn.make_text(self.src / "t.pdf", pages=4, label="Version A")
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        fake = CountingEngine("fake")
        w = self.worker({"fake": fake})
        jobs.create_job(con, [(did, "all")], method_override="fake")
        w.run_once()
        self.assertEqual(len(fake.calls), 4)
        # identical rerun: everything reused
        j2 = jobs.create_job(con, [(did, "all")], method_override="fake")
        w.run_once()
        self.assertEqual(len(fake.calls), 4)
        self.assertEqual(jobs.job_summary(con, j2)["pages_reused"], 4)
        # engine version change -> reprocess
        fake._version = "2"
        jobs.create_job(con, [(did, "all")], method_override="fake")
        w.run_once()
        self.assertEqual(len(fake.calls), 8)
        # source changed -> new SHA-256 -> nothing reused, warning shown
        time.sleep(0.01)
        syn.make_text(self.src / "t.pdf", pages=4, label="Version B")
        j4 = jobs.create_job(con, [(did, "all")], method_override="fake")
        w.run_once()
        s4 = jobs.job_summary(con, j4)
        self.assertEqual(len(fake.calls), 12)
        self.assertEqual(s4["pages_reused"], 0)
        self.assertIn("changed", s4["items"][0]["warning"])
        v = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        self.assertIn("Version B", next(r for r in v["results"] if r["engine"] == "textlayer")["text"])

    def test_ocr_settings_change_and_downstream_rules_no_rerun(self):
        syn.make_scanned(self.src / "s.pdf", pages=2)
        con = self.catalogue({"s.pdf": None})
        did = self.doc_id(con, "s.pdf")
        w = self.worker()
        tess = w.engines["tesseract"]
        calls = []
        orig = tess.extract_page
        tess.extract_page = lambda ctx: calls.append(ctx.page_no) or orig(ctx)
        job = jobs.create_job(con, [(did, "all")])
        w.run_once()
        self.assertEqual(calls, [1, 2])
        self.assertIn("SCANNED SHEET 2", jobs.page_view(con, did, 2, w.engines, w.rules.version)["results"][-1]["text"])
        # downstream rule change (dictionaries.json) + re-export: no OCR at all
        from pdf_pipeline.config import DEFAULT_CONFIG

        cfg = json.loads(DEFAULT_CONFIG.read_text())
        cfg["constant_columns"] = {"DATA_STORE_NAME": "Synthetic Store"}
        rules2 = self.base / "rules2.json"
        rules2.write_text(json.dumps(cfg))
        out = export_job(con, self.s, job, w.engines, w.rules.version, rules_path=rules2)
        self.assertEqual(calls, [1, 2])
        meta = list(csv.DictReader(open(out / "metadata.csv", encoding="utf-8-sig")))[0]
        self.assertEqual(meta["DATA_STORE_NAME"], "Synthetic Store")
        # OCR setting change -> new settings hash -> reprocessed
        tess.dpi = 150
        jobs.create_job(con, [(did, "all")])
        w.run_once()
        self.assertEqual(calls, [1, 2, 1, 2])

    def test_source_unavailable_then_recovers(self):
        syn.make_text(self.src / "t.pdf", pages=3)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        job = jobs.create_job(con, [(did, "all")])
        hidden = self.base / "nas_offline"
        self.src.rename(hidden)                               # simulate NAS disconnect
        w = self.worker()
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "waiting_for_source")
        hidden.rename(self.src)                               # NAS back
        con.execute("UPDATE jobs SET retry_after=0 WHERE id=?", (job,))
        w.run_once()
        self.assertEqual(jobs.job_summary(con, job)["status"], "done")

    def test_low_disk_pauses_and_keeps_work(self):
        syn.make_text(self.src / "t.pdf", pages=6)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        job = jobs.create_job(con, [(did, "all")])
        w = self.worker()
        real = shutil.disk_usage
        calls = {"n": 0}

        def fake_usage(p):
            calls["n"] += 1
            u = real(p)
            return type(u)(u.total, u.used, 0 if calls["n"] > 1 else u.free)

        w.s.min_free_gb = 1
        with mock.patch("pilot.worker.shutil.disk_usage", fake_usage):
            w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "paused")
        self.assertIn("disk", j["reason"])
        self.assertEqual(j["pages_done"], 3, "first batch kept")
        jobs.resume(con, job)
        w.run_once()
        self.assertEqual(jobs.job_summary(con, job)["status"], "done")

    def test_ocr_timeout_marks_page_failed(self):
        syn.make_scanned(self.src / "s.pdf", pages=2)
        con = self.catalogue({"s.pdf": None})
        did = self.doc_id(con, "s.pdf")
        w = self.worker()
        w.engines["tesseract"].timeout = 0.01
        job = jobs.create_job(con, [(did, "all")])
        w.run_once()
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done_with_errors")
        errs = [r[0] for r in con.execute("SELECT error FROM page_results WHERE engine='tesseract'")]
        self.assertEqual(errs, ["timeout", "timeout"])
        self.assertEqual(con.execute("SELECT COUNT(*) FROM page_results WHERE engine='textlayer' AND status='ok'")
                         .fetchone()[0], 2, "the text layer result is kept even when OCR fails")

    def test_encrypted_and_corrupt_do_not_stop_job(self):
        syn.make_encrypted(self.src / "locked.pdf")
        (self.src / "corrupt.pdf").write_bytes(b"%PDF-1.4 not really a pdf")
        syn.make_text(self.src / "good.pdf", pages=2)
        con = self.catalogue({"locked.pdf": None, "corrupt.pdf": None, "good.pdf": None})
        ids = [self.doc_id(con, n) for n in ("locked.pdf", "corrupt.pdf", "good.pdf")]
        job = jobs.create_job(con, [(i, "all") for i in ids])
        self.worker().run_once()
        items = {i["file_name"]: i for i in jobs.job_summary(con, job)["items"]}
        self.assertEqual(items["locked.pdf"]["status"], "failed")
        self.assertIn("password", items["locked.pdf"]["error"])
        self.assertEqual(items["corrupt.pdf"]["status"], "failed")
        self.assertEqual(items["good.pdf"]["status"], "done")
        enc = con.execute("SELECT encryption FROM sources WHERE doc_id=?", (ids[0],)).fetchone()[0]
        self.assertTrue(enc.startswith("Standard"))

    def test_page_range_beyond_end_fails_cleanly(self):
        syn.make_text(self.src / "t.pdf", pages=3)
        con = self.catalogue({"t.pdf": None})
        job = jobs.create_job(con, [(self.doc_id(con, "t.pdf"), "2-9")])
        self.worker().run_once()
        item = jobs.job_summary(con, job)["items"][0]
        self.assertEqual(item["status"], "failed")
        self.assertIn("beyond the last page", item["error"])

    def test_unimplemented_mode_rejected(self):
        syn.make_text(self.src / "t.pdf", pages=1)
        con = self.catalogue({"t.pdf": None})
        for mode in ("tables", "multimodal"):
            with self.assertRaises(jobs.JobError):
                jobs.create_job(con, [(self.doc_id(con, "t.pdf"), "all")], mode=mode)

    def test_review_revisions_keep_original(self):
        syn.make_text(self.src / "t.pdf", pages=2)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        job = jobs.create_job(con, [(did, "all")])
        w = self.worker()
        w.run_once()
        v = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        sha256 = v["sha256"]
        jobs.add_review(con, sha256, 1, v["results"][0]["id"], "approved")
        jobs.add_review(con, sha256, 1, v["results"][0]["id"], "corrected", "Corrected page one text")
        with self.assertRaises(jobs.JobError):
            jobs.add_review(con, sha256, 1, None, "corrected", "  ")
        v = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        self.assertEqual([r["status"] for r in v["review_history"]], ["approved", "corrected"])
        self.assertIn("Synthetic page 1", v["results"][0]["text"], "original extraction unchanged")
        out = export_job(con, self.s, job, w.engines, w.rules.version)
        pages = json.loads((out / f"doc{did}" / "pages.json").read_text())["pages"]
        self.assertIn("Synthetic page 1", pages[0]["original_text"])
        self.assertEqual(pages[0]["corrected_text"], "Corrected page one text")
        self.assertEqual(pages[1]["review_status"], "unreviewed")


# ---------------------------------------------------------------- crash recovery (real subprocess)
class KillResumeTest(Base):
    def test_sigkill_mid_job_then_resume(self):
        syn.make_scanned(self.src / "s.pdf", pages=12)
        con = self.catalogue({"s.pdf": None})
        did = self.doc_id(con, "s.pdf")
        job = jobs.create_job(con, [(did, "all")])
        cmd = [sys.executable, "-m", "pilot.worker", "--once", "--settings", str(self.settings_file)]
        proc = subprocess.Popen(cmd, cwd=PACKAGE_PARENT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        deadline = time.time() + 120
        while time.time() < deadline:
            n = con.execute("SELECT COUNT(*) FROM batches WHERE status='done'").fetchone()[0]
            if n >= 2:
                break
            time.sleep(0.2)
        os.kill(proc.pid, signal.SIGKILL)
        proc.wait()
        done_before = {(r[0], r[1]): r[2] for r in con.execute(
            "SELECT page_no, engine, id FROM page_results WHERE status='ok'")}
        self.assertGreaterEqual(len({p for p, _ in done_before}), 6)
        self.assertLess(len({p for p, e in done_before if e == "tesseract"}), 12, "killed before the end")
        subprocess.run(cmd, cwd=PACKAGE_PARENT, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=300)
        j = jobs.job_summary(con, job)
        self.assertEqual(j["status"], "done")
        rows = con.execute("SELECT page_no, engine, id FROM page_results WHERE status='ok'").fetchall()
        self.assertEqual(len(rows), 24, "12 pages x (textlayer + tesseract), no duplicates")
        after = {(r[0], r[1]): r[2] for r in rows}
        self.assertEqual({k: after[k] for k in done_before}, done_before, "completed pages were reused")
        # no half-written JSON files left behind
        self.assertEqual(list(self.s.results_dir.rglob(".tmp_*")), [])


# ---------------------------------------------------------------- web interface
class WebTest(Base):
    def test_api_and_guards(self):
        from fastapi.testclient import TestClient

        from pilot.web.app import create_app

        syn.make_mixed(self.src / "mixed.pdf")
        con = self.catalogue({"mixed.pdf": None})
        app = create_app(self.s, allowed_hosts={"testserver"})
        c = TestClient(app)
        tok = app.state.token
        r = c.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn(tok, r.text)
        self.assertNotRegex(r.text + (Path(__file__).parents[1] / "web/static/app.js").read_text(),
                            r"https?://(?!127\.0\.0\.1)", "no external resources")
        self.assertEqual(c.get("/docs").status_code, 404)
        self.assertEqual(c.get("/openapi.json").status_code, 404)
        self.assertEqual(c.get("/", headers={"host": "evil.example"}).status_code, 403)
        self.assertEqual(c.post("/api/jobs", json={}).status_code, 403, "token required")
        docs = c.get("/api/documents").json()
        did = docs[0]["id"]
        h = {"X-Pilot-Token": tok}
        self.assertEqual(c.post("/api/jobs", json={"items": [{"doc_id": did}], "mode": "tables"}, headers=h)
                         .status_code, 400)
        job = c.post("/api/jobs", json={"items": [{"doc_id": did, "pages": "1-3"}], "mode": "text_metadata"},
                     headers=h).json()["job_id"]
        self.worker().run_once()
        self.assertEqual(c.get(f"/api/jobs/{job}").json()["status"], "done")
        pv = c.get(f"/api/docs/{did}/pages/3").json()
        self.assertEqual(pv["selected_method"], "tesseract")
        img = c.get(f"/api/docs/{did}/pages/3/image")
        self.assertEqual(img.headers["content-type"], "image/png")
        self.assertTrue(img.content.startswith(b"\x89PNG"))
        rv = c.post(f"/api/docs/{did}/pages/3/review", json={"status": "approved"}, headers=h)
        self.assertEqual(rv.status_code, 200)
        rj = c.post(f"/api/docs/{did}/pages/3/retry", json={"method": "textlayer"}, headers=h).json()["job_id"]
        self.worker().run_once()
        self.assertEqual(c.get(f"/api/docs/{did}/pages/3").json()["override"], "textlayer")
        self.assertEqual(c.get(f"/api/jobs/{rj}").json()["status"], "done")
        report = c.get(f"/api/jobs/{job}/sanitized-report").text
        self.assertNotIn("mixed.pdf", report)
        self.assertNotIn(str(self.base), report)
        exp = c.post(f"/api/jobs/{job}/export", headers=h).json()["export_dir"]
        self.assertTrue((Path(exp) / "results.xlsx").is_file())

    def test_downloads(self):
        import io
        import zipfile

        from fastapi.testclient import TestClient

        from pilot.web.app import create_app

        syn.make_mixed(self.src / "mixed.pdf")
        con = self.catalogue({"mixed.pdf": None})
        did = self.doc_id(con, "mixed.pdf")
        job = jobs.create_job(con, [(did, "1-3")])
        w = self.worker()
        w.run_once()
        v = jobs.page_view(con, did, 1, w.engines, w.rules.version)
        jobs.add_review(con, v["sha256"], 1, None, "corrected", "Corrected cover text")
        app = create_app(self.s, allowed_hosts={"testserver"})
        c = TestClient(app)
        h = {"X-Pilot-Token": app.state.token}
        try:
            self.assertEqual(c.post(f"/api/docs/{did}/download", json={"format": "txt"}).status_code, 403, "token")
            r = c.post(f"/api/docs/{did}/download", json={"format": "txt", "pages": "1-2,5"}, headers=h)
            self.assertEqual(r.status_code, 200)
            self.assertIn('filename="mixed_pages_1-2_5.txt"', r.headers["content-disposition"])
            txt = r.text
            self.assertIn("===== Page 1 (method=textlayer; status=ok; review=corrected)", txt)
            self.assertIn("Corrected cover text", txt)
            self.assertIn("SCANNED PAGE TWO", txt)
            self.assertIn("===== Page 5", txt)
            self.assertIn("(not processed)", txt, "unprocessed pages are marked, not invented")
            rows = list(csv.DictReader(io.StringIO(c.post(f"/api/docs/{did}/download", headers=h,
                                                          json={"format": "csv", "pages": "1-2"}).content.decode("utf-8-sig"))))
            self.assertEqual([r["Page"] for r in rows], ["1", "2"])
            self.assertIn("FINAL SEISMIC PROCESSING REPORT", rows[0]["Original Text"], "original kept beside correction")
            self.assertEqual(rows[0]["Corrected Text"], "Corrected cover text")
            js = c.post(f"/api/docs/{did}/download", json={"format": "json", "pages": "3"}, headers=h).json()
            self.assertEqual(js["pages"][0]["method"], "tesseract")
            self.assertEqual(c.post(f"/api/docs/{did}/download", json={"format": "xml"}, headers=h).status_code, 400)
            self.assertEqual(c.post(f"/api/docs/{did}/download", json={"format": "txt", "pages": "99"}, headers=h)
                             .status_code, 400)
            z = c.post(f"/api/jobs/{job}/download", headers=h)
            self.assertEqual(z.headers["content-type"], "application/zip")
            names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
            self.assertTrue(any(n.endswith("/manifest.json") for n in names))
            self.assertTrue(any(n.endswith(f"/doc{did}/text.txt") for n in names))
        finally:
            app.state.source_monitor.stop()

    def test_hung_nas_does_not_freeze_the_app(self):
        """A share that never answers: every request must still return quickly."""
        from fastapi.testclient import TestClient

        from pilot.web import app as webapp

        syn.make_text(self.src / "t.pdf", pages=2)
        con = self.catalogue({"t.pdf": None})
        did = self.doc_id(con, "t.pdf")
        jobs.create_job(con, [(did, "all")])
        self.worker().run_once()
        release = __import__("threading").Event()
        hang = lambda *a, **k: release.wait(60)  # noqa: E731 - simulates an unresponsive NAS
        app = webapp.create_app(self.s, allowed_hosts={"testserver"}, source_probe=hang, nas_timeout=1.0)
        c = TestClient(app)
        try:
            t0 = time.time()
            self.assertEqual(c.get("/").status_code, 200)
            self.assertFalse(c.get("/api/state").json()["source_available"])
            self.assertLess(time.time() - t0, 2, "page and status answer without waiting for the NAS")
            with mock.patch.object(webapp, "stat_source", hang):
                t0 = time.time()
                for _ in range(3):  # repeated previews must not pile up
                    self.assertEqual(c.get(f"/api/docs/{did}/pages/1/image").status_code, 503)
                self.assertLess(time.time() - t0, 6)
            with mock.patch.object(webapp, "open_pdf", hang):  # NAS stalls during rendering
                r = c.get(f"/api/docs/{did}/pages/2/image")
                self.assertEqual(r.status_code, 503)
                r = c.get(f"/api/docs/{did}/pages/2/image")  # render still stuck -> quick 503, not a freeze
                self.assertEqual(r.status_code, 503)
            release.set()  # NAS answers again
            time.sleep(0.2)
            app.state.source_monitor.probe = lambda: True
            self.assertTrue(app.state.source_monitor.check_now())
            self.assertEqual(c.get(f"/api/docs/{did}/pages/2/image").status_code, 200, "recovers by itself")
        finally:
            release.set()
            app.state.source_monitor.stop()

    def test_no_telemetry_sdk_installed(self):
        import importlib.util

        self.assertIsNone(importlib.util.find_spec("opentelemetry.sdk"),
                          "only the no-op OpenTelemetry API should be present (pulled in by FastAPI)")


# ---------------------------------------------------------------- long document (slow)
@unittest.skipUnless(os.environ.get("RUN_SLOW"), "set RUN_SLOW=1 to run the 4,579-page test")
class LongDocumentTest(Base):
    def test_4579_pages(self):
        self.settings_file = make_env(self.base, batch_pages=25)
        self.s = load_settings(self.settings_file)
        t0 = time.time()
        syn.make_long_digital(self.src / "long.pdf", 4579)
        gen = time.time() - t0
        con = self.catalogue({"long.pdf": None})
        did = self.doc_id(con, "long.pdf")
        job = jobs.create_job(con, [(did, "all")])
        t0 = time.time()
        subprocess.run([sys.executable, "-m", "pilot.worker", "--once", "--settings", str(self.settings_file)],
                       cwd=PACKAGE_PARENT, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       timeout=3600)
        run = time.time() - t0
        rss_mb = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss / 1e6  # bytes on macOS
        j = jobs.job_summary(con, job)
        self.assertEqual((j["status"], j["pages_done"]), ("done", 4579))
        self.assertEqual(j["batches"], {"done": 184})
        w = self.worker()
        for n in (1, 25, 26, 2290, 4578, 4579):
            v = jobs.page_view(con, did, n, w.engines, w.rules.version)
            text = next(r for r in v["results"] if r["engine"] == "textlayer")["text"]
            self.assertIn(f"PAGE-MARKER {n}\n", text + "\n")
            self.assertIn(f"page {n} of 4579", text)
        size_mb = sum(p.stat().st_size for p in self.s.results_dir.rglob("*.json")) / 1e6
        print(f"\n[long-doc] generate {gen:.0f}s, process {run:.0f}s ({4579 / run:.0f} pages/s), "
              f"peak worker RSS {rss_mb:.0f} MB, page results {size_mb:.1f} MB")
        self.assertLess(rss_mb, 1500)


if __name__ == "__main__":
    unittest.main()
