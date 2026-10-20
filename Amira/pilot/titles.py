"""Catalogue titles from a second, read-only mapping workbook (e.g. Dulang_DIRt_v1.xlsx).

The title is taken from a configured column LETTER (default H), because in the DIRt sheets the
header above the titles is not "Title". With ``title_sheet = "*"`` every sheet that has the
location column (default "File Location") is read. Rows are joined to the inventory by the path
in that column (e.g. oss://bucket/20260925/DOCUMENTS/Seismic Report/X.pdf -> 20260925/DOCUMENTS/...),
never by file name alone. Blank titles stay blank; a path listed more than once (in any sheet)
gets no title. The title is catalogue metadata from that workbook - it is not extracted from the PDF.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl.utils import column_index_from_string

_SCHEME = re.compile(r"^[a-z][a-z0-9+.-]*://[^/]+/", re.IGNORECASE)  # oss://bucket/  s3://bucket/ ...


@dataclass
class TitleResult:
    titles: dict[str, str]                 # NFC relative path -> title (non-blank only)
    sources: dict[str, str] = field(default_factory=dict)   # relative path -> sheet name
    sheets: list[str] = field(default_factory=list)          # sheets that were read
    header_row: int = 0
    title_header: str = ""                 # header text above the title column in the first sheet read
    rows: int = 0
    blank: int = 0
    duplicates: list[str] = field(default_factory=list)
    bad_location: int = 0
    error: str = ""


def _nfc(s) -> str:
    return unicodedata.normalize("NFC", str(s).strip().replace("\\", "/"))


def location_to_rel(loc: str, path_prefix: str) -> str | None:
    """'oss://bucket/20260925/DOCUMENTS/X/a.pdf' -> '20260925/DOCUMENTS/X/a.pdf' (then minus path_prefix)."""
    loc = _nfc(loc)
    m = _SCHEME.match(loc)
    if m:
        loc = loc[m.end():]
    prefix = _nfc(path_prefix).strip("/")
    if not prefix:
        return loc.lstrip("/") or None
    prefix += "/"
    pos = loc.find(prefix)
    return loc[pos + len(prefix):] or None if pos >= 0 else None


def read_titles(workbook: Path, sheet: str, title_column: str, location_header: str, path_prefix: str,
                header_search_rows: int = 20) -> TitleResult:
    import openpyxl

    if not workbook.is_file():
        return TitleResult({}, error=f"title workbook not found: {workbook.name}")
    t_idx = column_index_from_string(title_column.strip().upper()) - 1
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)  # never saved
    try:
        names = wb.sheetnames if sheet.strip() == "*" else [sheet]
        if sheet.strip() != "*" and sheet not in wb.sheetnames:
            return TitleResult({}, error=f"worksheet '{sheet}' not found in title workbook")
        res = TitleResult({})
        seen: dict[str, int] = {}
        for name in names:
            it = wb[name].iter_rows(values_only=True)
            header, hrow = None, 0
            for i, r in enumerate(it, start=1):
                cells = [str(c).strip() if c is not None else "" for c in r]
                if location_header in cells:
                    header, hrow = cells, i
                    break
                if i >= header_search_rows:
                    break
            if header is None:
                continue  # not a mapping sheet (Read Me, Lookups, ...)
            if not res.sheets:
                res.header_row, res.title_header = hrow, header[t_idx] if t_idx < len(header) else ""
            res.sheets.append(name)
            loc_idx = header.index(location_header)
            for r in it:
                loc = r[loc_idx] if loc_idx < len(r) else None
                if loc in (None, ""):
                    continue
                res.rows += 1
                rel = location_to_rel(loc, path_prefix)
                if not rel:
                    res.bad_location += 1
                    continue
                seen[rel] = seen.get(rel, 0) + 1
                title = r[t_idx] if t_idx < len(r) else None
                title = " ".join(str(title).split()) if title not in (None, "") else ""
                if not title:
                    res.blank += 1
                elif seen[rel] == 1:
                    res.titles[rel] = title
                    res.sources[rel] = name
        if not res.sheets:
            res.error = f"no sheet with a '{location_header}' header"
        res.duplicates = sorted(k for k, n in seen.items() if n > 1)
        for k in res.duplicates:
            res.titles.pop(k, None)  # ambiguous: do not pick one
            res.sources.pop(k, None)
        return res
    finally:
        wb.close()


def load(settings) -> TitleResult:
    if not settings.title_workbook:
        return TitleResult({}, error="no title workbook configured")
    return read_titles(Path(settings.title_workbook).expanduser(), settings.title_sheet, settings.title_column,
                       settings.title_location_header, settings.path_prefix)
