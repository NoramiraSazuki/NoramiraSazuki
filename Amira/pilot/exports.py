"""Write local exports for a job into a NEW folder under <data_root>/exports (never overwrites).

    exports/job<id>_<timestamp>/
        manifest.json         engine versions, settings hashes, source SHA-256, page selections, caveats
        metadata.csv          one row per document (CLI schema), evidence.csv, pages.csv
        results.xlsx          the same three tables as sheets
        doc<id>/text.txt      selected text per original page (with page markers)
        doc<id>/pages.json    per page: original text, blocks, method, flags, review status, correction
        doc<id>/chunks.jsonl  paragraph chunks with provenance (for later vector search)
"""

from __future__ import annotations

import json
import time
from datetime import datetime
from pathlib import Path

from pdf_pipeline.config import DEFAULT_CONFIG, load_rules
from pdf_pipeline.export import CsvWriter
from pdf_pipeline.safety import create_new_dir
from pdf_pipeline.schema import EVIDENCE_COLUMNS, RESULT_COLUMNS

from . import __version__, derive
from .pages import parse_page_spec

PAGE_COLUMNS = ["Document ID", "File Name", "Catalogue Title", "Relative Path", "Page", "Method", "Auto Method", "Override",
                "Result Status", "Characters", "Validation Flags", "Warnings", "Review Status", "Has Correction",
                "Other Results", "Reading Order"]
CAVEAT = ("Processing status shows that an engine produced output; it does not mean the extraction is accurate. "
          "Only pages with review status 'approved' or 'corrected' were checked by a person.")


def _write_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    """Strings that start with '=' are stored as text, never as Excel formulas."""
    from openpyxl import Workbook

    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in rows:
            ws.append([None] * len(cols))
            for ci, c in enumerate(cols, start=1):
                v = r.get(c, "")
                cell = ws.cell(row=ws.max_row, column=ci)
                cell.value = v if not isinstance(v, str) else v[:32000]
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"
    with open(path, "xb") as fh:  # create only
        wb.save(fh)


def export_job(con, settings, job_id: int, engines: dict, routing_version: str,
               rules_path: Path = DEFAULT_CONFIG) -> Path:
    rules = load_rules(rules_path)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = create_new_dir(settings.exports_dir / f"job{job_id}_{stamp}")
    job = con.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
    items = con.execute(
        """SELECT i.*, d.rel_path, d.file_name, d.title AS catalogue_title, s.size, s.page_count AS src_pages, s.encryption, s.pdf_meta,
                  s.sha256 AS src_sha
           FROM job_items i JOIN documents d ON d.id=i.doc_id LEFT JOIN sources s ON s.doc_id=i.doc_id
           WHERE i.job_id=? ORDER BY i.id""", (job_id,)).fetchall()
    manifest = {
        "created": time.time(), "app_version": __version__, "job_id": job_id, "job_status": job["status"],
        "mode": job["mode"], "routing_version": routing_version, "caveat": CAVEAT,
        "engines": {n: {"version": e.version, "settings_hash": e.settings_hash, "settings": e.settings()}
                    for n, e in engines.items()},
        "extraction_rules": str(rules_path), "documents": [],
    }
    page_rows = []
    with CsvWriter(out / "metadata.csv", RESULT_COLUMNS) as meta_w, \
         CsvWriter(out / "evidence.csv", EVIDENCE_COLUMNS) as ev_w:
        meta_rows, ev_rows = [], []
        for it in items:
            entry = {"doc_id": it["doc_id"], "rel_path": it["rel_path"], "file_name": it["file_name"],
                     "catalogue_title": it["catalogue_title"],
                     "item_status": it["status"], "page_spec": it["page_spec"], "sha256": it["sha256"],
                     "warning": it["warning"], "error": it["error"]}
            manifest["documents"].append(entry)
            if not it["sha256"] or not it["page_count"]:
                continue
            if it["src_sha"] != it["sha256"]:
                entry["warning"] = ((entry["warning"] or "") + "; source changed after this job ran - export uses "
                                    "this job's version").strip("; ")
            pages = parse_page_spec(it["page_spec"], it["page_count"])
            sel = [derive.select_page(con, it["sha256"], p, engines, routing_version) for p in pages]
            doc = {"id": it["doc_id"], "rel_path": it["rel_path"], "file_name": it["file_name"],
                   "sha256": it["sha256"]}
            src = {"size": it["size"], "page_count": it["page_count"], "encryption": it["encryption"],
                   "pdf_meta": it["pdf_meta"]}
            row, evidence = derive.metadata_row(doc, src, sel, rules)
            meta_w.write(row)
            meta_rows.append(row)
            for e in evidence:
                ev_w.write(e)
            ev_rows.extend(evidence)

            ddir = out / f"doc{it['doc_id']}"
            ddir.mkdir()
            with open(ddir / "text.txt", "x", encoding="utf-8") as fh:
                for sp in sel:
                    tag = f"method={sp.method or '-'}; status={sp.status}; review={sp.review_status}"
                    fh.write(f"\n===== Page {sp.page_no} ({tag}) =====\n")
                    fh.write(sp.best_text)
                    fh.write("\n")
            with open(ddir / "pages.json", "x", encoding="utf-8") as fh:
                json.dump({"document": doc, "page_count": it["page_count"], "pages": [{
                    "page_no": sp.page_no, "method": sp.method, "auto_method": sp.auto_method,
                    "override": sp.override, "result_id": sp.result_id, "status": sp.status,
                    "original_text": sp.text, "review_status": sp.review_status,
                    "corrected_text": sp.corrected_text, "validation_flags": sp.flags, "warnings": sp.warnings,
                    "stats": sp.stats, "reading_order": sp.reading_order, "bbox_space": sp.bbox_space,
                    "blocks": sp.blocks, "other_results": sp.alternatives} for sp in sel]},
                    fh, ensure_ascii=False)
            with open(ddir / "chunks.jsonl", "x", encoding="utf-8") as fh:
                for sp in sel:
                    for ch in derive.chunks_for_page(doc, sp):
                        fh.write(json.dumps(ch, ensure_ascii=False) + "\n")
            for sp in sel:
                page_rows.append({
                    "Document ID": it["doc_id"], "File Name": it["file_name"],
                    "Catalogue Title": it["catalogue_title"] or "", "Relative Path": it["rel_path"],
                    "Page": sp.page_no, "Method": sp.method or "", "Auto Method": sp.auto_method or "",
                    "Override": sp.override or "", "Result Status": sp.status,
                    "Characters": sp.stats.get("chars", ""), "Validation Flags": "; ".join(sp.flags),
                    "Warnings": "; ".join(sp.warnings), "Review Status": sp.review_status,
                    "Has Correction": "yes" if sp.corrected_text else "", "Other Results": "; ".join(sp.alternatives),
                    "Reading Order": sp.reading_order or ""})
    with CsvWriter(out / "pages.csv", PAGE_COLUMNS) as pw:
        for r in page_rows:
            pw.write(r)
    _write_xlsx(out / "results.xlsx", {"metadata": (RESULT_COLUMNS, meta_rows),
                                       "evidence": (EVIDENCE_COLUMNS, ev_rows),
                                       "pages": (PAGE_COLUMNS, page_rows)})
    with open(out / "manifest.json", "x", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)
    return out
