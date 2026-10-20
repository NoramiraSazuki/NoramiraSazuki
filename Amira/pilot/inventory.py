"""Read the file-inventory workbook (read-only) and resolve each row to a file in the source folder.

Columns are found by header name, never by position or row number. Resolution never falls back
to matching by file name: a row resolves only when its exact relative path (after Unicode NFC
normalisation) exists under the source root. Anything else is reported.
"""

from __future__ import annotations

import os
import threading
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path, PurePosixPath

from pdf_pipeline.safety import is_within

REQUIRED_COLUMNS = ("File Path", "File Name")
# Optional inventory columns copied into the catalogue when present (shown, never changed).
OPTIONAL_COLUMNS = {
    "#": "inv_row_id",
    "Total Pages": "inv_pages",
    "Size (bytes)": "inv_size",
    "*Status of Documents/ Data Ingested": "inv_ingest_status",
    "Duplicate Group": "inv_duplicate_group",
    "Copies (same ETag)": "inv_copies",
    "ETag": "inv_etag",
}

# Resolution status codes
OK = "OK"
MISSING = "MISSING"
AMBIGUOUS = "AMBIGUOUS"
OUTSIDE_SOURCE = "OUTSIDE_SOURCE"
BAD_PATH = "BAD_PATH"
PREFIX_MISMATCH = "PREFIX_MISMATCH"
FOLDER_UNREADABLE = "FOLDER_UNREADABLE"   # its folder could not be listed (NAS error); not the same as missing
NOT_CHECKED = "NOT_CHECKED"               # listed in the inventory; not yet confirmed on the NAS


class InventoryError(RuntimeError):
    pass


@dataclass
class InventoryRow:
    rel_path: str                # relative to source_root, POSIX, NFC
    file_name: str
    resolve_status: str = ""
    resolved_path: str = ""
    note: str = ""
    extra: dict = field(default_factory=dict)  # optional inventory columns (inv_*)
    expected_path: str = ""      # source_root/rel_path, whether or not it was found (exact path, never by name)


@dataclass
class InventoryResult:
    rows: list[InventoryRow]
    unlisted: list[str]          # files on disk under source_root that no inventory row names
    source_available: bool
    header_row: int              # 1-based sheet row that held the header (informational)
    folders: list[str] = field(default_factory=list)  # every folder on disk + every inventory folder
    failed_folders: list[str] = field(default_factory=list)  # could not be listed (see FOLDER_UNREADABLE)


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text)


def _cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def read_inventory(workbook: Path, sheet: str, filter_column: str = "", filter_value: str = "",
                   extensions: tuple[str, ...] = ("pdf",), header_search_rows: int = 20) -> tuple[list[dict], int]:
    """Return (matching rows as dicts keyed by header, 1-based header row number).

    Rows are kept when their File Path ends with one of ``extensions`` and, if ``filter_column`` is
    set, that column equals ``filter_value``.
    """
    import openpyxl

    if not workbook.is_file():
        raise InventoryError(f"Inventory workbook not found: {workbook}")
    wanted = tuple("." + e.lower().lstrip(".") for e in extensions)
    need = REQUIRED_COLUMNS + ((filter_column,) if filter_column else ())
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)  # never saved
    try:
        if sheet not in wb.sheetnames:
            raise InventoryError(f"Worksheet '{sheet}' not found (have: {', '.join(wb.sheetnames)})")
        rows = wb[sheet].iter_rows(values_only=True)
        header, header_row = None, 0
        for i, r in enumerate(rows, start=1):
            names = [str(c).strip() if c is not None else "" for c in r]
            if all(c in names for c in need):
                header, header_row = names, i
                break
            if i >= header_search_rows:
                break
        if header is None:
            raise InventoryError(f"No header row with {', '.join(need)} in the first {header_search_rows} rows "
                                 f"of '{sheet}'")
        fi = header.index(filter_column) if filter_column else None
        pi = header.index("File Path")
        out = []
        for r in rows:
            if r is None or pi >= len(r) or r[pi] in (None, ""):
                continue
            if wanted and not str(r[pi]).strip().lower().endswith(wanted):
                continue
            if fi is not None and (fi >= len(r) or str(r[fi] or "").strip() != filter_value):
                continue
            out.append({h: _cell(v) for h, v in zip(header, r) if h})
        return out, header_row
    finally:
        wb.close()


@dataclass
class DiskIndex:
    files: dict[str, list[str]]          # NFC relative path -> actual relative paths (more than one = ambiguous)
    folders: list[str]                   # every visible folder (NFC, relative), including ones without PDFs
    failed: list[str] = field(default_factory=list)  # folders that could not be listed even after retries


def _list_dir(d: Path, timeout: float) -> list[tuple[str, str, bool]]:
    """List one folder in a helper thread; give up after ``timeout`` seconds.

    A stalled SMB request can block a thread in the kernel for many minutes and cannot be
    interrupted, so the scan moves on and the folder is reported instead of waiting.
    Returns (name, path, is_dir) tuples.
    """
    box: dict = {}

    def run():
        try:
            box["v"] = [(e.name, e.path, e.is_dir(follow_symlinks=False)) for e in os.scandir(d)]
        except OSError as exc:
            box["e"] = exc

    t = threading.Thread(target=run, daemon=True, name="scan-listdir")
    t.start()
    t.join(timeout)
    if t.is_alive():
        raise TimeoutError(f"folder listing took longer than {timeout:.0f}s")
    if "e" in box:
        raise box["e"]
    return box["v"]


def _scandir_retry(d: Path, attempts: int = 3, pause: float = 1.0, timeout: float = 60.0) -> list:
    """SMB shares occasionally fail or stall a listing; retry a few times before giving up."""
    for i in range(attempts):
        try:
            return _list_dir(d, timeout)
        except OSError:  # includes TimeoutError
            if i == attempts - 1:
                raise
            time.sleep(pause * (i + 1))
    return []


def _add_entries(idx: DiskIndex, source_root: Path, entries, wanted, stack=None) -> int:
    """Record one folder listing in the index. Returns the number of new matching files."""
    added = 0
    for name, path, is_dir in entries:
        if name.startswith("."):
            continue  # hidden files, macOS/NAS "._" companions
        rel = Path(path).relative_to(source_root).as_posix()
        if is_dir:
            if stack is not None:
                stack.append(Path(path))
            if _nfc(rel) not in idx.folders:
                idx.folders.append(_nfc(rel))
        elif name.lower().endswith(wanted):
            known = idx.files.setdefault(_nfc(rel), [])
            if rel not in known:  # a folder listed twice (retry/verification) adds nothing new
                known.append(rel)
                added += 1
            # No per-file stat here: over SMB that is one network round-trip per file (minutes for
            # thousands of PDFs). The worker checks size + SHA-256 before processing anyway.
    return added


def _disk_index(source_root: Path, extensions: tuple[str, ...] = ("pdf",), progress=None,
                folder_timeout: float = 60.0, on_listed=None) -> DiskIndex:
    """Walk the source tree (names only; file contents are never read).

    Folders that fail or stall are retried once more at the end (a second pass), because a
    network share often answers a little later. Only folders that fail twice are reported.
    """
    wanted = tuple("." + e.lower().lstrip(".") for e in extensions)
    idx = DiskIndex({}, [])
    counter = [0]

    def walk(stack: list[Path], failed: list[Path]) -> None:
        while stack:
            d = stack.pop()
            try:
                entries = _scandir_retry(d, timeout=folder_timeout)
            except OSError:
                failed.append(d)
                continue
            counter[0] += _add_entries(idx, source_root, entries, wanted, stack)
            if on_listed:  # this folder's files are now known: rows inside it can be checked
                on_listed(_nfc(d.relative_to(source_root).as_posix()) if d != source_root else "", idx)
            if progress:
                progress(len(idx.folders), counter[0])

    first_failed: list[Path] = []
    walk([source_root], first_failed)
    still_failed: list[Path] = []
    if first_failed:
        walk(list(reversed(first_failed)), still_failed)  # second pass
    idx.failed = sorted(_nfc(d.relative_to(source_root).as_posix()) if d != source_root else ""
                        for d in still_failed)
    idx.folders.sort()
    return idx


def _relative(file_path: str, prefix: str) -> tuple[str | None, str]:
    """Strip the inventory prefix and reject unsafe relative paths. Returns (rel, problem)."""
    fp = _nfc(str(file_path).strip().replace("\\", "/"))
    pre = _nfc(prefix.strip().replace("\\", "/"))
    if pre and not pre.endswith("/"):
        pre += "/"
    if not fp.startswith(pre):
        return None, PREFIX_MISMATCH
    rel = fp[len(pre):]
    parts = PurePosixPath(rel).parts
    if not rel or rel.startswith("/") or any(p in ("..", ".", "") for p in parts):
        return None, BAD_PATH
    return rel, ""


def prepare(rows: list[dict], source_root: Path, prefix: str) -> list[InventoryRow]:
    """Inventory rows without touching the NAS: path checks only; valid rows are NOT_CHECKED."""
    result = []
    for r in rows:
        extra = {col: r.get(src) for src, col in OPTIONAL_COLUMNS.items() if src in r}
        file_name = _nfc(str(r.get("File Name") or ""))
        rel, problem = _relative(r.get("File Path") or "", prefix)
        row = InventoryRow(rel_path=rel or _nfc(str(r.get("File Path") or "")), file_name=file_name, extra=extra)
        result.append(row)
        if problem:
            row.resolve_status = problem
            row.note = "inventory path does not start with the configured prefix" if problem == PREFIX_MISMATCH \
                else "unsafe relative path ('..', '.', or absolute)"
            continue
        row.expected_path = str(source_root / rel)
        row.resolve_status, row.note = NOT_CHECKED, "not yet confirmed on the NAS"
    # The same relative path listed more than once cannot be told apart: report, don't pick.
    counts: dict[str, int] = {}
    for row in result:
        counts[row.rel_path] = counts.get(row.rel_path, 0) + 1
    for row in result:
        if counts[row.rel_path] > 1:
            row.resolve_status, row.resolved_path = AMBIGUOUS, ""
            row.note = f"listed {counts[row.rel_path]} times in the inventory"
    return result


def inventory_folders(rows: list[InventoryRow]) -> set[str]:
    folders = set()
    for row in rows:
        parts = PurePosixPath(row.rel_path).parts[:-1]
        folders.update("/".join(parts[:i]) for i in range(1, len(parts) + 1))
    return folders


def _check(row: InventoryRow, disk: DiskIndex, source_root: Path, root_real: Path, final: bool) -> bool:
    """Resolve one NOT_CHECKED row against the files found so far. Returns False if it must wait."""
    rel = row.rel_path
    exact = disk.files.get(rel, [])
    if len(exact) == 1:
        candidate = source_root / exact[0]
        if not is_within(candidate, root_real):
            row.resolve_status, row.note = OUTSIDE_SOURCE, "resolves outside the source folder (symlink?)"
        else:
            row.resolve_status, row.resolved_path, row.note = OK, str(candidate), ""
            if PurePosixPath(rel).name != row.file_name:
                row.note = "File Name column differs from the name in File Path"
        return True
    if len(exact) > 1:
        row.resolve_status = AMBIGUOUS
        row.note = f"{len(exact)} files with Unicode-equivalent paths; not choosing one"
        return True
    if not final:
        return False
    if any(f == "" or rel.startswith(f + "/") for f in disk.failed):
        row.resolve_status = FOLDER_UNREADABLE
        row.note = "its folder could not be listed (NAS error) - not checked this time"
        return True
    folded = [a for k, actual in disk.files.items() if k.casefold() == rel.casefold() for a in actual]
    if folded:
        row.resolve_status = AMBIGUOUS
        row.note = f"only a case-different path exists ({len(folded)}); not matched automatically"
    else:
        row.resolve_status, row.note = MISSING, "no file at this relative path"
    return True


def check_on_disk(rows: list[InventoryRow], source_root: Path, extensions: tuple[str, ...] = ("pdf",),
                  progress=None, folder_timeout: float = 60.0, on_checked=None,
                  flush_seconds: float = 2.0) -> InventoryResult:
    """Walk the NAS and resolve NOT_CHECKED rows folder by folder.

    ``on_checked(rows)`` receives rows as soon as their folder has been listed (in small groups), so
    the catalogue fills in progressively; whatever is left is decided at the end.
    """
    available = source_root.is_dir()
    pending: dict[str, list[InventoryRow]] = {}
    for row in rows:
        if row.resolve_status == NOT_CHECKED:
            parent = row.rel_path.rsplit("/", 1)[0] if "/" in row.rel_path else ""
            pending.setdefault(parent, []).append(row)
    if not available:
        return InventoryResult(rows, [], False, 0, sorted(inventory_folders(rows)), [])
    root_real = source_root.resolve()
    done: list[InventoryRow] = []
    last = [time.monotonic()]

    def flush(force=False):
        if done and on_checked and (force or time.monotonic() - last[0] >= flush_seconds):
            on_checked(list(done))
            done.clear()
            last[0] = time.monotonic()

    def listed(dir_rel: str, disk: DiskIndex):
        for row in pending.pop(dir_rel, []):
            if _check(row, disk, source_root, root_real, final=False):
                done.append(row)
            else:
                pending.setdefault("\0final", []).append(row)  # not in its folder: decided at the end
        flush()

    disk = _disk_index(source_root, extensions, progress, folder_timeout, on_listed=listed)
    # Rows still undecided: their file was not in the listing, or their folder was never seen.
    # A network share can return an incomplete listing, so before calling anything MISSING the
    # folder is listed again directly; if that fails too, the row is FOLDER_UNREADABLE, not MISSING.
    wanted = tuple("." + e.lower().lstrip(".") for e in extensions)
    by_parent: dict[str, list[InventoryRow]] = {}
    for group in pending.values():
        for row in group:
            by_parent.setdefault(row.rel_path.rsplit("/", 1)[0] if "/" in row.rel_path else "", []).append(row)
    for parent, group in sorted(by_parent.items()):
        under_failed = any(f == "" or parent == f or parent.startswith(f + "/") for f in disk.failed)
        not_found = False
        if not under_failed:
            try:
                _add_entries(disk, source_root, _scandir_retry(source_root / parent, timeout=folder_timeout),
                             wanted)
            except FileNotFoundError:
                not_found = True
            except OSError:
                disk.failed.append(parent)
        for row in group:
            _check(row, disk, source_root, root_real, final=True)
            if not_found and row.resolve_status == MISSING:
                row.note = "folder not found on the NAS (confirmed by a second look-up)"
            done.append(row)
    flush(force=True)
    named = {r.rel_path for r in rows if r.resolve_status not in (BAD_PATH, PREFIX_MISMATCH)}
    unlisted = sorted(a for key, actual in disk.files.items() for a in actual if key not in named)
    folders = set(disk.folders) | inventory_folders(rows)
    return InventoryResult(rows, unlisted, True, 0, sorted(folders), sorted(disk.failed))


def resolve(rows: list[dict], source_root: Path, prefix: str, extensions: tuple[str, ...] = ("pdf",),
            progress=None, folder_timeout: float = 60.0) -> InventoryResult:
    """Prepare + check on disk in one go (all rows resolved when this returns)."""
    prepared = prepare(rows, source_root, prefix)
    res = check_on_disk(prepared, source_root, extensions, progress, folder_timeout)
    if not res.source_available:
        for row in prepared:
            if row.resolve_status == NOT_CHECKED:
                row.resolve_status, row.note = MISSING, "source folder not available (NAS not mounted?)"
    return res


def load(settings, progress=None) -> InventoryResult:
    rows, header_row = read_inventory(settings.inventory_workbook, settings.sheet, settings.filter_column,
                                      settings.filter_value, tuple(settings.extensions))
    res = resolve(rows, settings.source_root, settings.path_prefix, tuple(settings.extensions), progress,
                  settings.scan_folder_timeout)
    res.header_row = header_row
    return res
